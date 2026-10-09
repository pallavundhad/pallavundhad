"""Agent 1 - Reel Analyzer.

Takes a reel (local video file or a reel URL), pulls out the speech, hook,
on-screen captions and posted caption, studies how it is delivered, drafts new
captions, and writes everything into a Word document.

Pipeline: download (yt-dlp) -> audio + frames (ffmpeg) -> speech-to-text
(faster-whisper, runs locally) -> Claude reads frames + transcript -> .docx
"""

from __future__ import annotations

import base64
import datetime
import json
import os
import pathlib
import re
import subprocess
from typing import Any

from .base import structured

REPORTS = pathlib.Path(os.environ.get("HUB_REPORTS_DIR", "reports"))

ANALYSIS_SYSTEM = """You are an expert Instagram Reels analyst. You get frames from a reel (with timestamps),
its speech transcript and its posted caption. Break down exactly how the reel works:
- the hook in the first 3 seconds (spoken words, on-screen text, visual) and why it stops the scroll
- how the speech is delivered (tone, pace, energy, language style e.g. Hinglish, techniques like
  pattern interrupts, open loops, direct address, storytelling)
- every on-screen caption / text overlay you can read, with its timestamp
- the beat-by-beat structure
Then write 3 new draft captions in different styles and 5 alternative hooks the creator could use.
Be faithful to what is actually in the reel; if something is not visible or audible, say so."""

ANALYSIS_SCHEMA: dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "required": [
        "reel_summary", "hook", "speech_delivery", "on_screen_captions", "structure",
        "caption_breakdown", "hashtags_used", "draft_captions", "hook_variations", "takeaways",
    ],
    "properties": {
        "reel_summary": {"type": "string"},
        "hook": {
            "type": "object",
            "additionalProperties": False,
            "required": ["spoken", "on_screen_text", "visual", "hook_type", "why_it_works"],
            "properties": {k: {"type": "string"} for k in ["spoken", "on_screen_text", "visual", "hook_type", "why_it_works"]},
        },
        "speech_delivery": {
            "type": "object",
            "additionalProperties": False,
            "required": ["tone", "pace", "energy", "language_style", "techniques"],
            "properties": {
                "tone": {"type": "string"},
                "pace": {"type": "string"},
                "energy": {"type": "string"},
                "language_style": {"type": "string"},
                "techniques": {"type": "array", "items": {"type": "string"}},
            },
        },
        "on_screen_captions": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["time", "text"],
                "properties": {"time": {"type": "string"}, "text": {"type": "string"}},
            },
        },
        "structure": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["time", "beat"],
                "properties": {"time": {"type": "string"}, "beat": {"type": "string"}},
            },
        },
        "caption_breakdown": {"type": "string"},
        "hashtags_used": {"type": "array", "items": {"type": "string"}},
        "draft_captions": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["style", "caption"],
                "properties": {"style": {"type": "string"}, "caption": {"type": "string"}},
            },
        },
        "hook_variations": {"type": "array", "items": {"type": "string"}},
        "takeaways": {"type": "array", "items": {"type": "string"}},
    },
}


# ------------------------------ media steps --------------------------------

def fetch_reel(source: str, workdir: pathlib.Path) -> tuple[pathlib.Path, dict[str, Any]]:
    """Return (video path, metadata). Accepts a local file or a reel URL."""
    path = pathlib.Path(source).expanduser()
    if path.exists():
        return path, {"source": str(path)}

    import yt_dlp  # only needed for URLs

    opts: dict[str, Any] = {
        "outtmpl": str(workdir / "reel.%(ext)s"),
        "format": "mp4/best",
        "quiet": True,
        "no_warnings": True,
    }
    if os.environ.get("IG_COOKIES_FILE"):  # Instagram often needs a logged-in session
        opts["cookiefile"] = os.environ["IG_COOKIES_FILE"]
    with yt_dlp.YoutubeDL(opts) as ydl:
        info = ydl.extract_info(source, download=True)
        video = pathlib.Path(ydl.prepare_filename(info))
    meta = {
        "source": source,
        "creator": info.get("uploader") or info.get("channel"),
        "posted_caption": info.get("description") or "",
        "likes": info.get("like_count"),
        "comments": info.get("comment_count"),
        "views": info.get("view_count"),
        "upload_date": info.get("upload_date"),
    }
    return video, meta


def _duration(video: pathlib.Path) -> float:
    out = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(video)],
        capture_output=True, text=True, check=True,
    ).stdout.strip()
    return float(out or 0)


def extract_frames(video: pathlib.Path, workdir: pathlib.Path, max_frames: int = 16) -> list[tuple[float, pathlib.Path]]:
    """Dense frames in the first 3 s (the hook), then evenly spaced through the rest."""
    duration = _duration(video)
    hook_times = [t for t in (0.0, 0.5, 1.0, 1.5, 2.0, 3.0) if t < duration]
    rest = max_frames - len(hook_times)
    later = [3.0 + (duration - 3.0) * (i + 0.5) / rest for i in range(rest)] if duration > 4 else []
    frames = []
    for t in hook_times + later:
        out = workdir / f"frame_{t:07.2f}.jpg"
        subprocess.run(
            ["ffmpeg", "-y", "-loglevel", "error", "-ss", f"{t:.2f}", "-i", str(video),
             "-frames:v", "1", "-vf", "scale=540:-2", "-q:v", "4", str(out)],
            check=True,
        )
        if out.exists():
            frames.append((t, out))
    return frames


def transcribe(video: pathlib.Path, workdir: pathlib.Path) -> tuple[list[dict[str, Any]], str]:
    """Speech-to-text with timestamps. Returns (segments, note)."""
    audio = workdir / "audio.wav"
    result = subprocess.run(
        ["ffmpeg", "-y", "-loglevel", "error", "-i", str(video), "-vn", "-ac", "1", "-ar", "16000", str(audio)],
        capture_output=True, text=True,
    )
    if result.returncode != 0 or not audio.exists():
        return [], "No audio track found."
    try:
        from faster_whisper import WhisperModel
    except ImportError:
        return [], "Speech not transcribed: install faster-whisper (pip install faster-whisper)."
    model = WhisperModel(os.environ.get("WHISPER_MODEL", "small"), compute_type="int8")
    segments, info = model.transcribe(str(audio), vad_filter=True)
    segs = [{"start": round(s.start, 1), "end": round(s.end, 1), "text": s.text.strip()} for s in segments]
    note = f"Detected language: {info.language}" if segs else "No speech detected (music or silent reel)."
    return segs, note


# ------------------------------ word document ------------------------------

def write_docx(path: pathlib.Path, meta: dict[str, Any], analysis: dict[str, Any], segments: list[dict[str, Any]], speech_note: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Reel Breakdown", level=0)
    doc.add_paragraph(analysis["reel_summary"])

    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, key in [("Source", "source"), ("Creator", "creator"), ("Views", "views"),
                       ("Likes", "likes"), ("Comments", "comments"), ("Upload date", "upload_date")]:
        if meta.get(key) not in (None, ""):
            row = table.add_row().cells
            row[0].text, row[1].text = label, str(meta[key])

    doc.add_heading("1. Hook (first 3 seconds)", level=1)
    hook = analysis["hook"]
    for label, key in [("Spoken hook", "spoken"), ("On-screen text", "on_screen_text"), ("Visual", "visual"),
                       ("Hook type", "hook_type"), ("Why it works", "why_it_works")]:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(hook[key])

    doc.add_heading("2. Speech", level=1)
    sd = analysis["speech_delivery"]
    for label, key in [("Tone", "tone"), ("Pace", "pace"), ("Energy", "energy"), ("Language style", "language_style")]:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(sd[key])
    doc.add_paragraph("Delivery techniques:").runs[0].bold = True
    for t in sd["techniques"]:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_heading("Full transcript", level=2)
    doc.add_paragraph().add_run(speech_note).italic = True
    for s in segments:
        p = doc.add_paragraph()
        p.add_run(f"[{s['start']:.1f}s] ").bold = True
        p.add_run(s["text"])

    doc.add_heading("3. On-screen captions", level=1)
    if analysis["on_screen_captions"]:
        t = doc.add_table(rows=1, cols=2)
        t.style = "Light Grid Accent 1"
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Time", "Text on screen"
        for c in analysis["on_screen_captions"]:
            row = t.add_row().cells
            row[0].text, row[1].text = c["time"], c["text"]
    else:
        doc.add_paragraph("No on-screen text found.")

    doc.add_heading("4. Posted caption", level=1)
    doc.add_paragraph(meta.get("posted_caption") or "Not available (local file or caption not readable).")
    doc.add_paragraph(analysis["caption_breakdown"])
    if analysis["hashtags_used"]:
        doc.add_paragraph("Hashtags used: " + " ".join(analysis["hashtags_used"]))

    doc.add_heading("5. Structure, beat by beat", level=1)
    for b in analysis["structure"]:
        doc.add_paragraph(f"{b['time']} - {b['beat']}", style="List Bullet")

    doc.add_heading("6. Draft captions for you", level=1)
    for i, c in enumerate(analysis["draft_captions"], 1):
        doc.add_heading(f"Draft {i} - {c['style']}", level=2)
        doc.add_paragraph(c["caption"])

    doc.add_heading("7. Hook ideas you can use", level=1)
    for h in analysis["hook_variations"]:
        doc.add_paragraph(h, style="List Number")

    doc.add_heading("8. Key takeaways", level=1)
    for t in analysis["takeaways"]:
        doc.add_paragraph(t, style="List Bullet")

    doc.save(path)


# ------------------------------ entry point --------------------------------

def analyze_reel(source: str, notes: str = "") -> dict[str, Any]:
    """Run the full pipeline. Returns the analysis plus the path of the Word file."""
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    workdir = REPORTS / f"reel-{stamp}"
    workdir.mkdir(parents=True, exist_ok=True)

    video, meta = fetch_reel(source, workdir)
    frames = extract_frames(video, workdir)
    segments, speech_note = transcribe(video, workdir)

    content: list[dict[str, Any]] = []
    for t, frame in frames:
        content.append({"type": "text", "text": f"Frame at {t:.1f}s:"})
        content.append({
            "type": "image",
            "source": {"type": "base64", "media_type": "image/jpeg", "data": base64.standard_b64encode(frame.read_bytes()).decode()},
        })
    transcript = "\n".join(f"[{s['start']:.1f}-{s['end']:.1f}s] {s['text']}" for s in segments) or speech_note
    content.append({
        "type": "text",
        "text": (
            f"Reel metadata:\n{json.dumps(meta, ensure_ascii=False, indent=2)}\n\n"
            f"Speech transcript:\n{transcript}\n\n"
            + (f"Notes from the user: {notes}\n\n" if notes else "")
            + "Analyse this reel."
        ),
    })
    analysis = structured(ANALYSIS_SYSTEM, content, ANALYSIS_SCHEMA)

    name = re.sub(r"[^A-Za-z0-9]+", "-", (meta.get("creator") or video.stem))[:40].strip("-") or "reel"
    docx_path = REPORTS / f"reel-breakdown-{name}-{stamp}.docx"
    write_docx(docx_path, meta, analysis, segments, speech_note)
    return {"word_file": str(docx_path), "metadata": meta, "transcript": segments, "analysis": analysis}
